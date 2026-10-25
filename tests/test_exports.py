"""The export centre: three formats, and a file that matches the screen.

An export that disagrees with the grid it was started from is worse than no
export, so these tests read the produced files back rather than trusting that
a writer was called.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, datetime
from typing import ClassVar

import openpyxl
import pytest
from fastapi.testclient import TestClient
from fresh_schema import fresh_schema

from invoice_processing import exports
from invoice_processing.api import auth, service
from invoice_processing.api.service import app
from invoice_processing.db.models import ExportJob, Invoice, Tenant
from invoice_processing.db.repositories import (
    STALE_EXPORT_AFTER,
    PgExportJobs,
)
from invoice_processing.db.session import tenant_session

TENANT = "tenant-acme"
CLAIMS = {
    f"tok-{role}": {
        "user_id": f"u-{role}",
        "email": f"{role}@acme.test",
        "email_verified": True,
        "tenant_id": TENANT,
        "role": role,
    }
    for role in ("viewer", "clerk", "approver", "admin")
}


class FakeVerifier:
    def verify(self, token):
        if token not in CLAIMS:
            raise auth.AuthError("Invalid token")
        return CLAIMS[token]


class MemoryBlobs:
    """A blob store that keeps what it was given, so a file can be read back."""

    def __init__(self):
        self.stored: dict[str, bytes] = {}

    def put(self, key, data, content_type=None):
        self.stored[key] = data

    def get(self, key):
        if key not in self.stored:
            raise KeyError(key)
        return self.stored[key]


def approved(when="2026-08-15T09:00:00+00:00") -> dict:
    return {
        "decision": "approve",
        "decided_by": "approver@acme.test",
        "decided_at": when,
    }


@pytest.fixture
def db(scratch_db, monkeypatch):
    url = scratch_db("exports.db")
    monkeypatch.setenv("DATABASE_URL", url)
    monkeypatch.delenv("CLOUD_SQL_INSTANCE", raising=False)
    fresh_schema(url)
    with tenant_session(TENANT, url) as session:
        session.add(Tenant(id=TENANT, slug="acme", name="Acme"))
    monkeypatch.delenv("AUTH_DISABLED", raising=False)
    monkeypatch.setattr(auth, "_verifier", FakeVerifier())
    return url


@pytest.fixture
def blobs(monkeypatch):
    store = MemoryBlobs()
    monkeypatch.setattr(service, "get_blobs", lambda tenant_id: store)
    return store


@pytest.fixture
def client(db, blobs):
    def make(role="clerk"):
        return TestClient(app, headers={"Authorization": f"Bearer tok-{role}"})

    return make


def seed(db, count=3, **overrides):
    with tenant_session(TENANT, db) as session:
        for n in range(count):
            values = {
                "tenant_id": TENANT,
                "case_id": f"case-{n}",
                "vendor_name": f"Vendor {n}",
                "invoice_number": f"INV-{n}",
                "invoice_date": "2026-08-01",
                "currency": "AUD",
                "total_ex_tax": 1000 + n,
                "tax_amount": 100,
                "total_inc_tax": 1100 + n,
                "po_number": "PO-1",
                "min_confidence": 1,
                "fields": {"workflow": approved()},
            }
            values.update(overrides)
            session.add(Invoice(**values))


def export(client, **body):
    response = client().post("/api/sheet/exports", json=body)
    assert response.status_code == 200, response.text
    return response.json()


class TestWriters:
    """Each writer is checked by reading its output back."""

    COLUMNS: ClassVar[list[str]] = [
        "vendor_name",
        "invoice_number",
        "total_inc_tax",
        "quality_flags",
    ]
    LABELS: ClassVar[dict[str, str]] = {
        "vendor_name": "Vendor",
        "invoice_number": "Invoice no.",
        "total_inc_tax": "Total",
        "quality_flags": "Quality",
    }
    ROWS: ClassVar[list[dict]] = [
        {
            "vendor_name": "Café Ltd",
            "invoice_number": "INV-1",
            "total_inc_tax": 1100.5,
            "quality_flags": ["missing_po", "low_confidence"],
        },
        {
            "vendor_name": "Northwind",
            "invoice_number": None,
            "total_inc_tax": None,
            "quality_flags": [],
        },
    ]

    def test_csv_round_trips(self):
        raw = exports.write_csv(self.COLUMNS, self.ROWS, self.LABELS)
        parsed = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))

        assert parsed[0] == ["Vendor", "Invoice no.", "Total", "Quality"]
        assert parsed[1][0] == "Café Ltd"
        assert parsed[1][3] == "missing_po, low_confidence"
        # A missing value is an empty cell, not the word "None".
        assert parsed[2][1] == ""

    def test_csv_carries_a_bom_so_excel_reads_utf8(self):
        """Without it Excel renders "Café" as "CafÃ©", which reads as data
        corruption rather than an encoding default."""
        raw = exports.write_csv(self.COLUMNS, self.ROWS, self.LABELS)
        assert raw.startswith(b"\xef\xbb\xbf")

    def test_json_round_trips(self):
        payload = json.loads(
            exports.write_json(self.COLUMNS, self.ROWS, self.LABELS)
        )

        assert [c["key"] for c in payload["columns"]] == self.COLUMNS
        assert payload["rows"][0]["vendor_name"] == "Café Ltd"
        # A list stays a list in JSON: a machine reading this can use it.
        assert payload["rows"][0]["quality_flags"] == [
            "missing_po",
            "low_confidence",
        ]
        assert payload["exported_at"]

    def test_xlsx_round_trips(self):
        raw = exports.write_xlsx(self.COLUMNS, self.ROWS, self.LABELS)
        book = openpyxl.load_workbook(io.BytesIO(raw))
        sheet = book.active

        assert sheet.title == "Approved invoices"
        assert [c.value for c in sheet[1]] == [
            "Vendor",
            "Invoice no.",
            "Total",
            "Quality",
        ]
        assert sheet["A2"].value == "Café Ltd"

    def test_xlsx_totals_are_numbers_not_text(self):
        """So a finance user can sum the column instead of retyping it."""
        raw = exports.write_xlsx(self.COLUMNS, self.ROWS, self.LABELS)
        sheet = openpyxl.load_workbook(io.BytesIO(raw)).active

        assert isinstance(sheet["C2"].value, (int, float))
        assert sheet["C2"].value == 1100.5

    def test_extraction_text_cannot_become_a_formula(self):
        """Vendor names and descriptions are untrusted: "=HYPERLINK(...)"
        would run when the file is opened. Numbers stay numbers."""
        rows = [{"vendor_name": '=HYPERLINK("x")', "total_inc_tax": -85.0}]
        raw = exports.write_csv(self.COLUMNS, rows, self.LABELS)
        parsed = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        assert parsed[1][0] == '\'=HYPERLINK("x")'
        assert parsed[1][2] == "-85.0"

        sheet = openpyxl.load_workbook(
            io.BytesIO(exports.write_xlsx(self.COLUMNS, rows, self.LABELS))
        ).active
        assert sheet["A2"].value == '\'=HYPERLINK("x")'
        assert sheet["A2"].data_type == "s"
        assert sheet["C2"].value == -85.0

    def test_xlsx_freezes_the_header(self):
        raw = exports.write_xlsx(self.COLUMNS, self.ROWS, self.LABELS)
        sheet = openpyxl.load_workbook(io.BytesIO(raw)).active
        assert sheet.freeze_panes == "A2"

    def test_an_empty_export_is_still_a_valid_file(self):
        """Filters that match nothing produce headers, not a broken download."""
        for fmt, writer in exports.WRITERS.items():
            raw = writer(self.COLUMNS, [], self.LABELS)
            assert raw, fmt
        sheet = openpyxl.load_workbook(
            io.BytesIO(exports.write_xlsx(self.COLUMNS, [], self.LABELS))
        ).active
        assert [c.value for c in sheet[1]] == list(self.LABELS.values())


class TestColumnChoice:
    def test_no_columns_means_the_default_view(self):
        chosen = exports.resolve_columns(None)
        assert "vendor_name" in chosen
        assert "bank_account_number" not in chosen

    def test_an_unknown_column_is_dropped_not_fatal(self):
        """A view saved before a column was renamed still produces a file."""
        chosen = exports.resolve_columns(["vendor_name", "not_a_column"])
        assert chosen == ["vendor_name"]

    def test_only_unknown_columns_falls_back(self):
        assert exports.resolve_columns(["nope"]) == exports.resolve_columns(
            None
        )


class TestTheJob:
    def test_a_small_export_is_finished_when_the_request_returns(
        self, db, client, blobs
    ):
        """Waiting two seconds beats polling for them."""
        seed(db, 3)
        job = export(client, format="csv")

        assert job["status"] == "done"
        assert job["row_count"] == 3
        assert job["storage_key"] in blobs.stored

    def test_the_file_holds_the_rows_the_grid_would_show(
        self, db, client, blobs
    ):
        seed(db, 2)
        job = export(client, format="csv", columns=["invoice_number"])

        raw = blobs.stored[job["storage_key"]].decode("utf-8-sig")
        numbers = [row[0] for row in list(csv.reader(io.StringIO(raw)))[1:]]
        assert sorted(numbers) == ["INV-0", "INV-1"]

    def test_filters_narrow_the_export(self, db, client, blobs):
        seed(db, 3)
        with tenant_session(db and TENANT, db) as session:
            session.add(
                Invoice(
                    tenant_id=TENANT,
                    case_id="other",
                    vendor_name="Excluded",
                    invoice_number="INV-X",
                    currency="INR",
                    total_inc_tax=5,
                    fields={"workflow": approved()},
                )
            )

        job = export(client, format="csv", filters={"currency": "INR"})
        assert job["row_count"] == 1

    def test_selected_rows_win_over_filters(self, db, client, blobs):
        """A ticked selection is the request; re-running the filter could
        return something else if an invoice was approved in between."""
        seed(db, 3)
        job = export(client, format="csv", case_ids=["case-1"])

        assert job["row_count"] == 1
        raw = blobs.stored[job["storage_key"]].decode("utf-8-sig")
        assert "Vendor 1" in raw
        assert "Vendor 0" not in raw

    def test_an_unapproved_invoice_is_not_exported(self, db, client):
        """The export inherits the sheet's definition of approved."""
        seed(db, 1)
        with tenant_session(TENANT, db) as session:
            session.add(
                Invoice(
                    tenant_id=TENANT,
                    case_id="pending",
                    vendor_name="Not approved",
                    invoice_number="INV-P",
                    fields={},
                )
            )

        job = export(client, format="csv")
        assert job["row_count"] == 1

    def test_an_unknown_format_is_refused_before_a_job_exists(self, db, client):
        seed(db, 1)
        response = client().post("/api/sheet/exports", json={"format": "pdf"})

        assert response.status_code == 400
        assert "pdf" in response.text
        assert client().get("/api/sheet/exports").json()["exports"] == []

    def test_a_failure_is_recorded_on_the_job(self, db, client, monkeypatch):
        """The caller has already been answered, so the job is the only place
        left to say what went wrong."""
        seed(db, 1)

        def explode(*args, **kwargs):
            raise RuntimeError("disk is full")

        monkeypatch.setitem(exports.WRITERS, "csv", explode)
        response = client().post("/api/sheet/exports", json={"format": "csv"})

        assert response.status_code == 200
        job = response.json()
        assert job["status"] == "failed"
        assert "disk is full" in job["error"]

    def test_jobs_are_listed_newest_first(self, db, client):
        seed(db, 1)
        export(client, format="csv")
        export(client, format="json")

        listed = client().get("/api/sheet/exports").json()["exports"]
        assert [job["format"] for job in listed] == ["json", "csv"]


class TestDownload:
    def test_a_finished_export_downloads_with_a_filename(self, db, client):
        seed(db, 1)
        job = export(client, format="xlsx")

        response = client().get(f"/api/sheet/exports/{job['id']}/download")
        assert response.status_code == 200
        assert "attachment" in response.headers["content-disposition"]
        assert ".xlsx" in response.headers["content-disposition"]
        # Really a spreadsheet, not an error page with a nice header.
        assert openpyxl.load_workbook(io.BytesIO(response.content)).active

    def test_an_unfinished_export_says_so(self, db, client):
        seed(db, 1)
        job = PgExportJobs(db).create(
            TENANT, requested_by="clerk@acme.test", fmt="csv", spec={}
        )

        response = client().get(f"/api/sheet/exports/{job['id']}/download")
        assert response.status_code == 409
        assert "queued" in response.text

    def test_an_unknown_export_is_404(self, db, client):
        assert (
            client().get("/api/sheet/exports/nope/download").status_code == 404
        )


class TestRoles:
    def test_a_viewer_cannot_request_an_export(self, db, client):
        seed(db, 1)
        assert (
            client("viewer")
            .post("/api/sheet/exports", json={"format": "csv"})
            .status_code
            == 403
        )

    def test_a_viewer_can_download_one(self, db, client):
        """Reading the ledger and taking a copy of what you can already read
        are the same permission."""
        seed(db, 1)
        job = export(client, format="csv")

        assert (
            client("viewer")
            .get(f"/api/sheet/exports/{job['id']}/download")
            .status_code
            == 200
        )


class TestAudit:
    def test_requesting_and_finishing_are_both_recorded(
        self, db, client, monkeypatch
    ):
        recorded = []

        class RecordingAudit:
            def emit(self, ctx, event, payload, before=None):
                recorded.append(event)

        monkeypatch.setattr(service, "_audit", RecordingAudit())
        seed(db, 1)
        export(client, format="csv")

        assert "export_requested" in recorded
        assert "export_completed" in recorded


class TestAStuckJobCanBeRestarted:
    """The runner lives in the API process, so a deploy mid-export orphans the
    job. It has to be visible as stuck and startable again, or the user's only
    recourse is to rebuild the filters and hope."""

    def test_a_fresh_job_is_not_stale(self, db, client):
        seed(db, 1)
        job = export(client, format="csv")
        assert job["stale"] is False

    def test_an_unfinished_job_goes_stale(self, db, client):
        seed(db, 1)
        job = export(client, format="csv")
        long_ago = datetime.now(UTC) - STALE_EXPORT_AFTER * 2
        with tenant_session(TENANT, db) as session:
            row = session.get(ExportJob, job["id"])
            row.status = "running"
            row.created_at = long_ago

        assert PgExportJobs().get(TENANT, job["id"])["stale"] is True

    def test_retrying_a_failed_job_produces_the_file(
        self, db, client, blobs, monkeypatch
    ):
        seed(db, 2)

        def explode(*args, **kwargs):
            raise RuntimeError("disk is full")

        working = exports.WRITERS["csv"]
        monkeypatch.setitem(exports.WRITERS, "csv", explode)
        job = client().post("/api/sheet/exports", json={"format": "csv"}).json()
        assert job["status"] == "failed"

        monkeypatch.setitem(exports.WRITERS, "csv", working)
        retried = client().post(f"/api/sheet/exports/{job['id']}/retry")
        assert retried.status_code == 200, retried.text

        # The background task runs inline under TestClient, so by now it is done.
        finished = client().get(f"/api/sheet/exports/{job['id']}").json()
        assert finished["status"] == "done"
        assert finished["row_count"] == 2
        assert finished["error"] is None

    def test_a_finished_job_cannot_be_retried(self, db, client):
        seed(db, 1)
        job = export(client, format="csv")

        response = client().post(f"/api/sheet/exports/{job['id']}/retry")
        assert response.status_code == 409

    def test_retrying_needs_clerk(self, db, client):
        seed(db, 1)
        job = export(client, format="csv")

        response = client("viewer").post(
            f"/api/sheet/exports/{job['id']}/retry"
        )
        assert response.status_code == 403


def _csv_cells(blobs, job):
    raw = blobs.stored[job["storage_key"]].decode("utf-8-sig")
    return list(csv.reader(io.StringIO(raw)))[1:]


class TestBankNumbersMaskByRole:
    NUMBER = "901031925572"

    def test_a_clerk_export_carries_the_last_four(self, db, client, blobs):
        seed(db, 1, bank_account_number=self.NUMBER)
        job = export(
            lambda: client("clerk"),
            format="csv",
            columns=["bank_account_number"],
        )
        assert _csv_cells(blobs, job) == [["…5572"]]

    def test_an_approver_export_carries_the_number(self, db, client, blobs):
        seed(db, 1, bank_account_number=self.NUMBER)
        job = export(
            lambda: client("approver"),
            format="csv",
            columns=["bank_account_number"],
        )
        assert _csv_cells(blobs, job) == [[self.NUMBER]]

    def test_a_job_without_the_flag_masks(self, db, client, blobs):
        """Jobs queued before masking existed say nothing about the caller;
        the safe reading is the restrictive one."""

        seed(db, 1, bank_account_number=self.NUMBER)
        job = PgExportJobs().create(
            TENANT,
            requested_by="approver@acme.test",
            fmt="csv",
            spec={
                "columns": ["bank_account_number"],
                "filters": {},
                "case_ids": [],
            },
        )
        exports.run(job["id"], TENANT)
        done = PgExportJobs().get(TENANT, job["id"])
        assert _csv_cells(blobs, done) == [["…5572"]]

    def test_a_download_is_not_marked_no_store(self, db, client):
        seed(db, 1)
        job = export(client, format="csv")
        response = client().get(f"/api/sheet/exports/{job['id']}/download")
        assert "no-store" not in response.headers.get("cache-control", "")
